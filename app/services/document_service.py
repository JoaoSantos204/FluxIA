import os
import logging
from pathlib import Path
from pypdf import PdfReader
from docx import Document

logger = logging.getLogger(__name__)


def ler_documento(caminho_arquivo: str | Path) -> str:
    caminho = Path(caminho_arquivo)
    extensao = caminho.suffix.lower()

    if extensao == ".txt":
        return ler_txt(caminho)
    elif extensao == ".pdf":
        return ler_pdf(caminho)
    elif extensao == ".docx":
        return ler_docx(caminho)
    elif extensao == ".pptx":
        return ler_pptx(caminho)
    elif extensao in {".png", ".jpg", ".jpeg"}:
        return ler_imagem(caminho)
    else:
        raise ValueError(
            f"Tipo de documento não suportado: {extensao}. Tipos permitidos: .pdf, .docx, .txt, .pptx, .png, .jpg, .jpeg"
        )


def ler_txt(caminho: Path) -> str:
    with open(caminho, "r", encoding="utf-8", errors="ignore") as arquivo:
        return arquivo.read()


def ler_pdf(caminho: Path) -> str:
    leitor = PdfReader(caminho)
    texto = ""
    for pagina in leitor.pages:
        texto_pagina = pagina.extract_text()
        if texto_pagina:
            texto += texto_pagina + "\n"
    return texto


def ler_docx(caminho: Path) -> str:
    documento = Document(caminho)
    texto = ""
    for paragrafo in documento.paragraphs:
        if paragrafo.text.strip():
            texto += paragrafo.text + "\n"
    return texto


def ler_pptx(caminho: Path) -> str:
    """Extrai texto estruturado de apresentações PowerPoint (.pptx)."""
    from pptx import Presentation

    try:
        prs = Presentation(caminho)
    except Exception as e:
        logger.error(f"[DocumentService] Erro ao abrir PPTX {caminho.name}: {e}")
        raise ValueError(f"Não foi possível abrir o arquivo PowerPoint (.pptx): {e}")

    slides_texto = []
    for num, slide in enumerate(prs.slides, 1):
        linhas_slide = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        linhas_slide.append(t)
        if linhas_slide:
            slides_texto.append(f"--- Slide {num} ---\n" + "\n".join(linhas_slide))

    return "\n\n".join(slides_texto)


def ler_imagem(caminho: Path) -> str:
    """
    Utiliza o modelo multimodal Gemini para descrever detalhadamente e transcrever
    o conteúdo textual (OCR) e visual de imagens (.png, .jpg, .jpeg).
    """
    from google import genai
    from google.genai import types

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY não configurada para leitura multimodal de imagens.")

    client = genai.Client(api_key=api_key)
    ext = caminho.suffix.lower()
    mime = "image/png" if ext == ".png" else "image/jpeg"

    with open(caminho, "rb") as f:
        image_bytes = f.read()

    part = types.Part.from_bytes(data=image_bytes, mime_type=mime)
    prompt = (
        "Você é um especialista em OCR e análise documental para bases de conhecimento corporativas. "
        "Analise minuciosamente a imagem a seguir e extraia: "
        "1. Todo o texto visível com exatidão (transcrição completa, títulos, tabelas, dados). "
        "2. Uma descrição detalhada do conteúdo visual, contexto e informações relevantes. "
        "Retorne todo esse conteúdo em Português do Brasil para que seja vetorizado em um banco semântico."
    )

    try:
        response = client.models.generate_content(
            model="gemini-3.5-flash-lite",
            contents=[part, prompt]
        )
        texto = response.text or ""
        if not texto.strip():
            response_fb = client.models.generate_content(
                model="gemini-3.6-flash",
                contents=[part, prompt]
            )
            texto = response_fb.text or ""

        if not texto.strip():
            raise ValueError("A IA não conseguiu extrair texto ou descrição útil desta imagem.")

        return texto.strip()

    except Exception as e:
        logger.error(f"[DocumentService] Erro ao extrair texto da imagem {caminho.name}: {e}")
        raise ValueError(f"Não foi possível processar a imagem com IA: {e}")
